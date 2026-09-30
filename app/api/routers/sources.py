from __future__ import annotations

import csv
import io
import json
import uuid
from pathlib import Path

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query

from app.api.auth import get_current_user
from app.autofill.source_reader import read_source

router = APIRouter(prefix="/sources", tags=["sources"])

_SOURCES_BASE = Path("outputs/sources")
_META_BASE = Path("outputs/sources/meta")

_ASIN_KEYS = {
    "asin", "amazon_asin", "asin_number", "amazon_seller_sku",
    "amz_sku", "amazon_sku", "asin_code", "amazon_id",
    "item_id", "itemid", "product_id", "productid",
}
_SKU_KEYS = {"sku", "supplier_part_number", "part_number", "item_number"}
_CAT_KEYS = {"category", "category_name", "product_category", "template_category"}


def _norm_key(k: str) -> str:
    import re
    k = k.strip().lower()
    k = re.sub(r"[\s\-/]+", "_", k)
    k = re.sub(r"[^\w]", "", k)
    return k


def _detect_special_cols(columns: list[str]) -> tuple[str | None, str | None, str | None]:
    norm_cols = {_norm_key(c): c for c in columns}
    asin_col = next((norm_cols[k] for k in _ASIN_KEYS if k in norm_cols), None)
    sku_col = next((norm_cols[k] for k in _SKU_KEYS if k in norm_cols), None)
    cat_col = next((norm_cols[k] for k in _CAT_KEYS if k in norm_cols), None)
    return asin_col, sku_col, cat_col


def _sources_dir(user_id: str) -> Path:
    return _SOURCES_BASE / user_id


def _meta_dir(user_id: str) -> Path:
    return _META_BASE / user_id


def _save_meta(user_id: str, source_id: str, meta: dict) -> None:
    md = _meta_dir(user_id)
    md.mkdir(parents=True, exist_ok=True)
    (md / f"{source_id}.json").write_text(
        json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8"
    )


def _load_meta(user_id: str, source_id: str) -> dict | None:
    p = _meta_dir(user_id) / f"{source_id}.json"
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


@router.post("/csv/upload")
async def upload_csv(file: UploadFile = File(...), current_user: dict = Depends(get_current_user)):
    sd = _sources_dir(current_user["user_id"])
    sd.mkdir(parents=True, exist_ok=True)
    source_id = str(uuid.uuid4())
    dest = sd / f"{source_id}_{file.filename}"
    content = await file.read()
    dest.write_bytes(content)

    try:
        rows = read_source(dest)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Cannot parse CSV: {exc}")

    columns = list(rows[0].keys()) if rows else []
    asin_col, sku_col, cat_col = _detect_special_cols(columns)

    meta = {
        "source_id": source_id,
        "filename": file.filename,
        "stored_path": str(dest),
        "row_count": len(rows),
        "columns": columns,
        "asin_column": asin_col,
        "sku_column": sku_col,
        "category_column": cat_col,
    }
    _save_meta(current_user["user_id"], source_id, meta)

    return {
        "source_id": source_id,
        "filename": file.filename,
        "row_count": len(rows),
        "columns": columns,
        "asin_column": asin_col,
        "sku_column": sku_col,
        "category_column": cat_col,
        "preview": rows[:5],
    }


@router.post("/excel/upload")
async def upload_excel(
    file: UploadFile = File(...),
    sheet: str | None = Query(default=None),
    current_user: dict = Depends(get_current_user),
):
    sd = _sources_dir(current_user["user_id"])
    sd.mkdir(parents=True, exist_ok=True)
    source_id = str(uuid.uuid4())
    dest = sd / f"{source_id}_{file.filename}"
    content = await file.read()
    dest.write_bytes(content)

    try:
        wb = openpyxl.load_workbook(str(dest), data_only=True)
        sheet_names = wb.sheetnames
        chosen_sheet = sheet if (sheet and sheet in sheet_names) else sheet_names[0]
        ws = wb[chosen_sheet]

        headers = []
        for c in range(1, (ws.max_column or 1) + 1):
            v = ws.cell(1, c).value
            headers.append(str(v).strip() if v is not None else f"col_{c}")

        rows = []
        for r in range(2, min((ws.max_row or 1) + 1, 10002)):
            row = {}
            for c, h in enumerate(headers, 1):
                v = ws.cell(r, c).value
                row[h] = str(v).strip() if v is not None else ""
            rows.append(row)
        wb.close()
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Cannot parse Excel: {exc}")

    norm_headers = [_norm_key(h) for h in headers]
    asin_col, sku_col, cat_col = _detect_special_cols(headers)

    meta = {
        "source_id": source_id,
        "filename": file.filename,
        "stored_path": str(dest),
        "row_count": len(rows),
        "columns": norm_headers,
        "asin_column": asin_col,
        "sku_column": sku_col,
        "category_column": cat_col,
        "sheet_names": sheet_names,
        "chosen_sheet": chosen_sheet,
    }
    _save_meta(current_user["user_id"], source_id, meta)

    return {
        "source_id": source_id,
        "filename": file.filename,
        "sheet_names": sheet_names,
        "chosen_sheet": chosen_sheet,
        "row_count": len(rows),
        "columns": norm_headers,
        "asin_column": asin_col,
        "sku_column": sku_col,
        "category_column": cat_col,
        "preview": rows[:5],
    }


@router.get("/{source_id}/preview")
async def preview_source(source_id: str, current_user: dict = Depends(get_current_user)):
    meta = _load_meta(current_user["user_id"], source_id)
    if not meta:
        raise HTTPException(status_code=404, detail=f"Source {source_id!r} not found")

    stored_path = Path(meta["stored_path"])
    if not stored_path.exists():
        raise HTTPException(status_code=404, detail="Source file not found on disk")

    try:
        rows = read_source(stored_path)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Cannot read source: {exc}")

    return {
        "source_id": source_id,
        "filename": meta["filename"],
        "row_count": len(rows),
        "columns": meta.get("columns", []),
        "asin_column": meta.get("asin_column"),
        "sku_column": meta.get("sku_column"),
        "preview": rows[:20],
    }
