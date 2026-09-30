from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Any

import openpyxl
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse

from app.analyser.template_registry import TemplateRegistry
from app.autofill.source_reader import read_source
from app.api.auth import get_current_user
from app.api.session_store import run_store
from app.api import profile_store
import logging

logger = logging.getLogger(__name__)
from app.api.models import (
    CreateRunResponse,
    ResolveRequest,
    ResolvedProductOut,
    PreviewRowOut,
    PreviewCellOut,
    ManualEditsRequest,
    ValidationFindingOut,
    FindingsGrouped,
    ExportResponse,
    DownloadFileOut,
)

router = APIRouter(prefix="/runs", tags=["runs"])

# Brand/company names to strip from Product Name regardless of source data.
# Longer variants must come before shorter ones so they match first.
_PRODUCT_NAME_STRIP_TERMS: list[str] = [
    "DC Voltage",
    "DCVoltage",
    "LED Sone",
    "LEDSone",
    "SRM",
]

_REGISTRY_BASE = Path("outputs/template_registry")
_PROFILE_BASE = Path("outputs/template_profiles")


def _registry_json(user_id: str) -> Path:
    return _REGISTRY_BASE / user_id / "template_registry.json"


def _profile_dir(user_id: str) -> Path:
    return _PROFILE_BASE / user_id
_SOURCES_META_BASE = Path("outputs/sources/meta")
_SOURCES_BASE = Path("outputs/sources")


def _sources_meta_dir(user_id: str) -> Path:
    return _SOURCES_META_BASE / user_id
_CONFIG_DIR = Path("config")

# ── helpers ───────────────────────────────────────────────────────────────────

def _load_registry(user_id: str) -> TemplateRegistry:
    return TemplateRegistry.load(_registry_json(user_id))


def _norm(k: str) -> str:
    k = k.strip().lower()
    k = re.sub(r"[\s\-/]+", "_", k)
    k = re.sub(r"[^\w]", "", k)
    return k


def _load_source_meta(user_id: str, source_id: str) -> dict | None:
    p = _sources_meta_dir(user_id) / f"{source_id}.json"
    if not p.exists():
        return None
    return json.loads(p.read_text(encoding="utf-8"))


def _load_source_rows(user_id: str, source_id: str) -> list[dict]:
    meta = _load_source_meta(user_id, source_id)
    if not meta:
        return []
    try:
        return read_source(Path(meta["stored_path"]))
    except Exception:
        return []


def _load_field_mappings() -> list[dict]:
    import yaml
    mp = _CONFIG_DIR / "field_mappings.yaml"
    if not mp.exists():
        return []
    with open(mp, encoding="utf-8") as f:
        data = yaml.safe_load(f)
    return data.get("mappings", [])


def _load_fixed_values() -> list[dict]:
    """Return raw fixed-value entries (list of dicts with wayfair_field, value, optional categories)."""
    import yaml
    fp = _CONFIG_DIR / "fixed_values.yaml"
    if not fp.exists():
        return []
    with open(fp, encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    entries = data.get("fixed_values", [])
    if isinstance(entries, list):
        return entries
    # legacy dict format — convert to list with no category restriction
    if isinstance(entries, dict):
        return [{"wayfair_field": k, "value": v} for k, v in entries.items()]
    return []


def _fixed_values_for_category(entries: list[dict], category: str) -> dict[str, str]:
    """Filter fixed-value entries by category and return {wayfair_field: value}."""
    result: dict[str, str] = {}
    cat_norm = category.strip().lower()
    for e in entries:
        cats = e.get("categories")
        if cats:
            # Category-specific: only apply if category matches (case-insensitive substring)
            if not any(cat_norm in c.lower() or c.lower() in cat_norm for c in cats):
                continue
        field = e.get("wayfair_field", "")
        value = e.get("value", "")
        if field:
            result[field] = value
    return result


def _load_profile(user_id: str, template_id: str, inline_profiles: dict | None = None) -> dict | None:
    pd = _profile_dir(user_id)
    # 1. Try disk first (fast path after normal upload)
    registry = _load_registry(user_id)
    entry = registry.find_by_id(template_id)
    if entry:
        candidates = [
            Path(entry.profile_path),
            Path("outputs") / entry.profile_path,
            pd / Path(entry.profile_path).name,
        ]
        for candidate in candidates:
            if candidate.exists():
                try:
                    return json.loads(candidate.read_text(encoding="utf-8"))
                except Exception:
                    pass
                break

    # 1b. Check for a profile cached by template_id on disk (written during resolve)
    cached_by_id = pd / f"{template_id}.json"
    if cached_by_id.exists():
        try:
            return json.loads(cached_by_id.read_text(encoding="utf-8"))
        except Exception:
            pass

    # 1c. Check PostgreSQL — persists across Railway container restarts
    db_profile = profile_store.load_profile(user_id, template_id)
    if db_profile:
        # Re-populate disk cache for fast subsequent reads
        try:
            pd.mkdir(parents=True, exist_ok=True)
            (pd / f"{template_id}.json").write_text(json.dumps(db_profile), encoding="utf-8")
        except Exception:
            pass
        return db_profile

    # 2. Fall back to inline profile sent by the browser
    if inline_profiles and template_id in inline_profiles:
        profile = inline_profiles[template_id]
        if isinstance(profile, dict):
            # Persist to PostgreSQL (survives restarts) and disk (fast reads)
            profile_store.save_profile(user_id, template_id, profile)
            try:
                pd.mkdir(parents=True, exist_ok=True)
                cache_path = (
                    pd / Path(entry.profile_path).name
                    if entry
                    else pd / f"{template_id}.json"
                )
                cache_path.parent.mkdir(parents=True, exist_ok=True)
                cache_path.write_text(json.dumps(profile), encoding="utf-8")
            except Exception:
                pass
            return profile

    return None


def _apply_transform(value: str, transform: str) -> str:
    """Apply a named transform rule to a value."""
    if not transform:
        return value
    if transform.startswith("prefix:"):
        prefix = transform[len("prefix:"):]
        return value if value.startswith(prefix) else f"{prefix}{value}"
    if transform.startswith("suffix:"):
        suffix = transform[len("suffix:"):]
        return value if value.endswith(suffix) else f"{value}{suffix}"
    if transform == "upper":
        return value.upper()
    if transform == "lower":
        return value.lower()
    if transform == "title_case":
        return value.title()
    if transform.startswith("truncate:"):
        try:
            n = int(transform.split(":")[1])
            return value[:n]
        except ValueError:
            pass
    if transform == "marketing_copy":
        import re as _re
        # Strip trademark, copyright, registered and other special symbols
        value = _re.sub(r'[™©®℠†‡•·°]', '', value)
        # Strip Unicode mathematical bold/italic/script letters (𝐀-𝟿)
        value = _re.sub(r'[\U0001D400-\U0001D7FF]', '', value)
        # Strip HTML tags if any
        value = _re.sub(r'<[^>]+>', ' ', value)
        # Collapse newlines, tabs and multiple spaces into a single space
        value = _re.sub(r'[\r\n\t]+', ' ', value)
        value = _re.sub(r' {2,}', ' ', value)
        value = value.strip()
        # Truncate to Wayfair limit
        return value[:4000]
    return value


def _map_row_to_template(
    src_row: dict[str, str],
    exact_headers: list[str],
    mappings: list[dict],
    fixed_values: dict[str, str],
    manual_edits: dict[str, str],
    asin: str,
    category: str,
) -> dict[str, PreviewCellOut]:
    cells: dict[str, PreviewCellOut] = {}

    # Build alias -> wayfair_field lookup and transform map
    alias_map: dict[str, str] = {}
    required_fields: set[str] = set()
    transform_map: dict[str, str] = {}   # wayfair_field -> transform rule
    for m in mappings:
        wf = m.get("wayfair_field", "")
        cats = m.get("categories", ["all"])
        if "all" not in cats and category not in cats:
            continue
        for alias in m.get("aliases", []):
            alias_map[_norm(alias)] = wf
        if m.get("required"):
            required_fields.add(wf)
        if m.get("transform"):
            transform_map[wf] = m["transform"]

    # Build normalized source row
    norm_src = {_norm(k): v for k, v in src_row.items()}

    for header in exact_headers:
        edit_key = f"{asin}::{header}"
        if edit_key in manual_edits:
            cells[header] = PreviewCellOut(value=manual_edits[edit_key], status="manual")
            continue

        norm_header = _norm(header)

        # Check fixed value first
        if norm_header in {_norm(k): k for k in fixed_values}:
            fv_key = next((k for k in fixed_values if _norm(k) == norm_header), None)
            if fv_key:
                cells[header] = PreviewCellOut(value=fixed_values[fv_key], status="fixed_value")
                continue

        # Try alias mapping
        matched_alias = None
        for alias, wf in alias_map.items():
            if wf == header and alias in norm_src:
                matched_alias = alias
                break

        raw_val = ""
        if matched_alias and norm_src.get(matched_alias, ""):
            raw_val = norm_src[matched_alias]
        else:
            # Fallback: direct normalised column name match
            raw_val = norm_src.get(norm_header, "")

        if raw_val:
            final_val = _apply_transform(raw_val, transform_map.get(header, ""))
            if norm_header == "product_name":
                import re as _re
                # Strip dynamic brand from source row
                brand = ""
                for b_alias in ("brand", "manufacturer", "brand_name"):
                    if norm_src.get(b_alias, "").strip():
                        brand = norm_src[b_alias].strip()
                        break
                if brand:
                    final_val = _re.sub(_re.escape(brand) + r"[\s\-–—,]*", "", final_val, flags=_re.IGNORECASE).strip()
                # Strip known company/brand terms regardless of source data
                for term in _PRODUCT_NAME_STRIP_TERMS:
                    final_val = _re.sub(r'(?<!\w)' + _re.escape(term) + r'[\s\-–—,]*', "", final_val, flags=_re.IGNORECASE).strip()
            cells[header] = PreviewCellOut(value=final_val, status="source_value")
        elif header in required_fields:
            cells[header] = PreviewCellOut(value="", status="missing")
        else:
            cells[header] = PreviewCellOut(value="", status="unmapped")

    # ── Conditional derived fields ────────────────────────────────────────────
    _POWER_SOURCE_CATEGORIES = {
        "pendant lights", "flush mount lighting", "table lamps", "wall sconces",
        "chandeliers", "lighting accessories", "transformers and switches dimmers outlets",
    }
    cat_norm = category.strip().lower()
    if any(c in cat_norm for c in _POWER_SOURCE_CATEGORIES) or any(cat_norm in c for c in _POWER_SOURCE_CATEGORIES):
        sku = norm_src.get("supplier_part_number", "") or norm_src.get("sku", "") or ""
        title = norm_src.get("product_name", "") or norm_src.get("title", "") or norm_src.get("name", "") or ""
        description = norm_src.get("description", "") or norm_src.get("marketing_copy", "") or norm_src.get("product_description", "") or ""
        bullets = " ".join(
            norm_src.get(f"bullet_{i}", "") or norm_src.get(f"feature_{i}", "") or norm_src.get(f"feature_bullet_{i}", "") or ""
            for i in range(1, 6)
        )
        combined = " ".join([sku, title, description, bullets]).lower()
        is_plugin = "plug-in" in combined or "plugin" in combined

        derived = {
            "Power Source":  "Plug-in" if is_plugin else "Hardwired",
            "Plug Type":     "Type G: 220–240 Volt" if is_plugin else "Unavailable",
            "Adapter Type":  "G - BS 1363-3 Compliant" if is_plugin else "Does Not Apply",
        }
        for field, value in derived.items():
            if field in exact_headers and f"{asin}::{field}" not in manual_edits:
                cells[field] = PreviewCellOut(value=value, status="fixed_value")

    # ── Bulb Replaceability ───────────────────────────────────────────────────
    _BULB_CATEGORIES = {
        "pendant lights", "flush mount lighting", "table lamps", "wall sconces", "chandeliers",
    }
    if any(c in cat_norm for c in _BULB_CATEGORIES) or any(cat_norm in c for c in _BULB_CATEGORIES):
        if "Bulb Replaceability" in exact_headers and f"{asin}::Bulb Replaceability" not in manual_edits:
            bulb_included_val = cells.get("Bulb Included", PreviewCellOut(value="", status="unmapped")).value.strip().lower()
            if bulb_included_val == "yes":
                bulb_replaceability = "Yes"
            else:
                bulb_replaceability = "Does Not Apply"
            cells["Bulb Replaceability"] = PreviewCellOut(value=bulb_replaceability, status="fixed_value")

    return cells


def _expand_parent_asin(asin: str) -> list[str]:
    """If asin is a parentAsin, return its child ASINs. Otherwise return [asin]."""
    try:
        from app.sources.configurator_sot_source import resolve_asin
        identity = resolve_asin(asin)
        if identity.status == "parent_asin" and identity.child_asins:
            logger.info("Expanded parent ASIN %s → %d children", asin, len(identity.child_asins))
            return identity.child_asins
    except Exception:
        pass
    return [asin]


def _get_sot_row(asin: str) -> dict[str, str]:
    """Resolve ASIN via SOT and return a flat dict keyed by Wayfair header names."""
    try:
        import yaml
        from app.sources.configurator_sot_source import (
            resolve_asin, get_sot_attributes, get_image_urls, get_title, get_bullet_points,
        )
        from app.sources.sot_value_parser import parse_sot_value, convert_unit
    except Exception:
        return {}

    sot_mappings_path = _CONFIG_DIR / "configurator_sot_mappings.yaml"
    if not sot_mappings_path.exists():
        return {}
    with open(sot_mappings_path, encoding="utf-8") as f:
        sot_mappings = yaml.safe_load(f).get("mappings", [])
    approved = [m for m in sot_mappings if m.get("mapping_status") == "approved"]

    identity = resolve_asin(asin)
    if identity.status != "resolved":
        logger.info("SOT: ASIN %s not resolved (%s)", asin, identity.status)
        return {}

    sku = identity.primary_sku

    _SYNTHETIC = frozenset({
        "__sku__", "__image_url_1__", "__image_url_2__", "__image_url_3__",
        "__image_url_4__", "__image_url_5__", "__video_url_1__",
        "__bullet_1__", "__bullet_2__", "__bullet_3__", "__bullet_4__", "__bullet_5__",
        "__description__", "__base_cost__",
    })
    needed_keys = [m["sot_key"] for m in approved if m["sot_key"] not in _SYNTHETIC]
    attrs = get_sot_attributes(sku, needed_keys or None)
    attr_by_key: dict[str, str] = {}
    for a in attrs:
        if a.attribute_key not in attr_by_key:
            attr_by_key[a.attribute_key] = a.value

    image_urls: list[str] | None = None
    bullet_points: list[str] | None = None

    def _imgs() -> list[str]:
        nonlocal image_urls
        if image_urls is None:
            image_urls = get_image_urls(sku)
        return image_urls

    def _bullets() -> list[str]:
        nonlocal bullet_points
        if bullet_points is None:
            bullet_points = get_bullet_points(sku, asin=asin)
        return bullet_points

    def _base_cost() -> tuple[str | None, bool]:
        """Returns (value, unverified). unverified=True when price came from wrong_sku=1 row."""
        from app.sources.postgres_connection import readonly_cursor
        try:
            with readonly_cursor() as cur:
                cur.execute("""
                    SELECT COALESCE(price, 0), COALESCE(price_per_order, 0)
                    FROM public.listing_data
                    WHERE ref_id = %s AND wrong_sku = 0
                      AND (status ILIKE 'active' OR status IN ('DISCOVERABLE', 'BUYABLE', '1'))
                      AND which_channel_name ILIKE '%%amazon%%'
                      AND price IS NOT NULL
                    ORDER BY (CASE WHEN market_place = 'UK' THEN 0 ELSE 1 END), id
                    LIMIT 1
                """, (asin,))
                row = cur.fetchone()
            if row:
                return str(round((float(row[0]) + float(row[1])) * 0.60, 2)), False
            # Fallback: wrong_sku=1
            with readonly_cursor() as cur:
                cur.execute("""
                    SELECT COALESCE(price, 0), COALESCE(price_per_order, 0)
                    FROM public.listing_data
                    WHERE ref_id = %s AND (status ILIKE 'active' OR status IN ('DISCOVERABLE', 'BUYABLE', '1'))
                      AND which_channel_name ILIKE '%%amazon%%'
                      AND price IS NOT NULL
                    ORDER BY (CASE WHEN market_place = 'UK' THEN 0 ELSE 1 END), id
                    LIMIT 1
                """, (asin,))
                row = cur.fetchone()
            if row:
                return str(round((float(row[0]) + float(row[1])) * 0.60, 2)), True
        except Exception:
            pass
        return None, False

    result: dict[str, str] = {}
    for mapping in approved:
        sot_key = mapping["sot_key"]
        header = mapping["wayfair_header"]
        transform = mapping.get("transformation")

        if sot_key == "__sku__":
            raw = sku
        elif sot_key.startswith("__image_url_"):
            idx = int(sot_key.replace("__image_url_", "").replace("__", "")) - 1
            urls = _imgs()
            raw = urls[idx] if idx < len(urls) else None
        elif sot_key == "__video_url_1__":
            raw = None
        elif sot_key == "__description__":
            from app.sources.configurator_sot_source import get_description as _get_desc
            raw = _get_desc(sku, asin=asin)
        elif sot_key.startswith("__bullet_"):
            idx = int(sot_key.replace("__bullet_", "").replace("__", "")) - 1
            bps = _bullets()
            raw = bps[idx] if idx < len(bps) else None
        elif sot_key == "__base_cost__":
            raw, _bc_unverified = _base_cost()
            if raw and _bc_unverified:
                result["__base_cost_unverified__"] = "true"
        else:
            raw = attr_by_key.get(sot_key)

        if raw is None:
            continue

        parsed = parse_sot_value(raw)
        val = raw
        if transform == "mm_to_cm":
            if parsed.value_type == "measurement" and parsed.numeric is not None:
                val = str(convert_unit(parsed, "cm").parsed_value)
            else:
                try:
                    import re as _re2
                    n = float(_re2.match(r"^([\d.]+)", raw.strip()).group(1))
                    val = str(round(n * 0.1, 4))
                except Exception:
                    pass
        elif transform == "g_to_kg":
            if parsed.value_type == "measurement" and parsed.numeric is not None:
                val = str(convert_unit(parsed, "kg").parsed_value)
        elif transform == "mm_to_numeric":
            import re as _re2
            m2 = _re2.match(r"^([\d.]+)", raw.strip())
            val = m2.group(1) if m2 else raw
        elif transform == "marketing_copy":
            import re as _re2
            val = _re2.sub(r'[™©®℠†‡•·°]', '', val)
            val = _re2.sub(r'[\U0001D400-\U0001D7FF]', '', val)
            val = _re2.sub(r'<[^>]+>', ' ', val)
            val = _re2.sub(r'[\r\n\t]+', ' ', val)
            val = _re2.sub(r' {2,}', ' ', val).strip()[:4000]

        result[header] = val

    return result


# ── routes ────────────────────────────────────────────────────────────────────

@router.post("", response_model=CreateRunResponse)
async def create_run(current_user: dict = Depends(get_current_user)):
    run_id = run_store.create_run(current_user["user_id"])
    return CreateRunResponse(run_id=run_id)


@router.get("/{run_id}")
async def get_run(run_id: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")
    return state


@router.post("/{run_id}/resolve")
async def resolve_run(run_id: str, body: ResolveRequest, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    registry = _load_registry(current_user["user_id"])
    use_sot = "configurator_sot" in body.source_ids

    inline_profiles = body.template_profiles or {}
    raw_asins = [a.strip().upper() for a in body.asins if a.strip()]
    normalized_asins: list[str] = raw_asins

    # Load source rows from all source_ids (skip the virtual SOT source)
    all_source_rows: list[tuple[str, dict]] = []  # (source_id, row)
    for src_id in body.source_ids:
        if src_id == "configurator_sot":
            continue
        meta = _load_source_meta(current_user["user_id"], src_id)
        if not meta:
            continue
        rows = _load_source_rows(current_user["user_id"], src_id)
        asin_col = meta.get("asin_column")
        sku_col = meta.get("sku_column")
        for row in rows:
            all_source_rows.append((src_id, row))

    # Build ASIN/SKU index using detected columns from metadata
    asin_row_index: dict[str, list[tuple[str, dict]]] = {}
    _ASIN_NORMS = {
        "asin", "amazon_asin", "asin_number", "amazon_seller_sku",
        "amz_sku", "amazon_sku", "asin_code", "amazon_id",
        "item_id", "itemid", "product_id", "productid",
    }
    _SKU_NORMS = {"sku", "supplier_part_number", "part_number", "item_number", "style_number"}

    for src_id, row in all_source_rows:
        meta = _load_source_meta(current_user["user_id"], src_id)
        detected_asin_col = meta.get("asin_column") if meta else None
        detected_sku_col = meta.get("sku_column") if meta else None

        for col, val in row.items():
            val_clean = str(val).strip().upper()
            if not val_clean:
                continue
            norm_col = _norm(col)
            # Use detected column name first, fall back to known patterns
            is_asin_col = (col == detected_asin_col) or (norm_col in _ASIN_NORMS)
            is_sku_col = (col == detected_sku_col) or (norm_col in _SKU_NORMS)
            if is_asin_col or is_sku_col:
                asin_row_index.setdefault(val_clean, []).append((src_id, row))

    # Cache any inline profiles to disk so preview_run / export_run can load them
    # without needing inline_profiles (works even when the registry has no entry).
    for tid, prof in inline_profiles.items():
        _load_profile(current_user["user_id"], tid, inline_profiles)

    # Determine template routing — prefer registry entries; fall back to lightweight
    # stubs built from inline_profiles so template_id is never stored as null.
    templates = [registry.find_by_id(tid) for tid in body.template_ids if registry.find_by_id(tid)]
    if not templates and inline_profiles:
        # Build minimal stub objects from inline profiles so downstream code
        # can set template_id / template_filename on each resolved product.
        class _InlineTemplateStub:
            __slots__ = ("template_id", "filename", "category")
            def __init__(self, tid: str, prof: dict) -> None:
                self.template_id = tid
                self.filename = Path(prof.get("source_file", tid)).name
                self.category = prof.get("category_name", "")
        for tid in body.template_ids:
            if tid in inline_profiles:
                templates.append(_InlineTemplateStub(tid, inline_profiles[tid]))
    single_template = templates[0] if len(templates) == 1 else None

    # Batch-fetch all SOT rows — also expands parent ASINs to children
    sot_rows_by_asin: dict[str, dict] = {}
    if use_sot:
        try:
            import yaml
            from app.sources.configurator_sot_source import batch_get_sot_rows
            sot_mappings_path = _CONFIG_DIR / "configurator_sot_mappings.yaml"
            sot_mappings: list[dict] = []
            if sot_mappings_path.exists():
                with open(sot_mappings_path, encoding="utf-8") as _f:
                    sot_mappings = yaml.safe_load(_f).get("mappings", [])
            sot_rows_by_asin, normalized_asins = batch_get_sot_rows(normalized_asins, sot_mappings)
        except Exception as exc:
            logger.error("batch_get_sot_rows failed: %s", exc)

    resolved_products = []
    for asin in normalized_asins:
        matches = asin_row_index.get(asin, [])
        sot_row = sot_rows_by_asin.get(asin, {}) if use_sot else {}

        if not matches:
            if use_sot:
                tpl = single_template or (templates[0] if templates else None)
                resolved_products.append({
                    "asin": asin,
                    "sku": sot_row.get("Supplier Part Number"),
                    "product_name": sot_row.get("Product Name"),
                    "template_id": tpl.template_id if tpl else None,
                    "template_filename": tpl.filename if tpl else None,
                    "source_id": "configurator_sot",
                    "source_row": sot_row,
                    "status": "matched" if sot_row else "unmatched",
                })
            else:
                resolved_products.append({
                    "asin": asin,
                    "sku": None,
                    "product_name": None,
                    "template_id": single_template.template_id if single_template else (templates[0].template_id if templates else None),
                    "template_filename": single_template.filename if single_template else (templates[0].filename if templates else None),
                    "source_id": None,
                    "source_row": {},
                    "status": "unmatched",
                })
            continue

        src_id, row = matches[0]
        sku = row.get("sku") or row.get("supplier_part_number") or None
        product_name = row.get("product_name") or row.get("title") or row.get("name") or None

        # Merge SOT data on top of CSV row (SOT fills gaps, CSV takes priority)
        if use_sot and sot_row:
            merged = {**sot_row, **{k: v for k, v in row.items() if v}}
            row = merged
            sku = sku or sot_row.get("Supplier Part Number")
            product_name = product_name or sot_row.get("Product Name")

        # Route to template
        if single_template:
            tpl = single_template
        else:
            cat_val = row.get("category_name") or row.get("category") or ""
            tpl = next(
                (t for t in templates if t.category.lower() in cat_val.lower() or cat_val.lower() in t.category.lower()),
                templates[0] if templates else None,
            )

        resolved_products.append({
            "asin": asin,
            "sku": sku,
            "product_name": product_name,
            "template_id": tpl.template_id if tpl else None,
            "template_filename": tpl.filename if tpl else None,
            "source_id": src_id,
            "source_row": row,
            "status": "matched",
        })

    # Apply Brand/MPN fixed values at resolve time so preview shows correct values
    # without requiring a separate sot-fill step.
    resolve_edits: dict[str, str] = {}
    if use_sot:
        for prod in resolved_products:
            _asin = prod.get("asin", "")
            _sot = sot_rows_by_asin.get(_asin, {})
            if not _sot:
                continue
            _sku = _sot.get("Supplier Part Number") or prod.get("sku", "")
            if _sot.get("__is_us__") == "true":
                resolve_edits[f"{_asin}::Brand"] = "Neighbour Market"
                resolve_edits[f"{_asin}::Manufacturer Part Number"] = f"NM-{_sku}"
                # Blank out UK document fields for US ASINs
                for _i in range(1, 4):
                    for _f in (
                        f"Document File Name or URL {_i}",
                        f"Document Type {_i}",
                        f"Legal Document Type {_i}",
                        f"Document Region {_i}",
                    ):
                        resolve_edits[f"{_asin}::{_f}"] = ""
            else:
                resolve_edits[f"{_asin}::Brand"] = "LEDSone"
                resolve_edits[f"{_asin}::Manufacturer Part Number"] = f"WF-{_sku}"

    run_store.update_run(
        current_user["user_id"],
        run_id,
        status="resolved",
        uploaded_templates=body.template_ids,
        source_files=body.source_ids,
        asins=body.asins,
        resolved_products=resolved_products,
        cached_profiles=inline_profiles,  # stored so preview/export never need to re-fetch
        **({"manual_edits": resolve_edits} if resolve_edits else {}),
    )

    return {"run_id": run_id, "resolved_count": len(resolved_products), "products": resolved_products}


@router.get("/{run_id}/debug-source")
async def debug_source_row(run_id: str, asin: str, current_user: dict = Depends(get_current_user)):
    """Debug: show what keys are in the resolved source_row for an ASIN."""
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail="Run not found")
    products = state.get("resolved_products", [])
    for prod in products:
        if prod.get("asin") == asin.upper():
            src = prod.get("source_row", {})
            bullet_keys = {k: v for k, v in src.items() if "bullet" in k.lower() or "feature" in k.lower()}
            return {
                "asin": asin,
                "status": prod.get("status"),
                "source_id": prod.get("source_id"),
                "all_keys": list(src.keys()),
                "bullet_keys": bullet_keys,
            }
    raise HTTPException(status_code=404, detail="ASIN not found in run")


@router.get("/{run_id}/preview")
async def preview_run(run_id: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    products = state.get("resolved_products", [])
    manual_edits = state.get("manual_edits", {})
    cached_profiles = state.get("cached_profiles", {})
    mappings = _load_field_mappings()
    fv_entries = _load_fixed_values()

    rows = []
    for prod in products:
        template_id = prod.get("template_id")
        if not template_id:
            continue
        profile = _load_profile(current_user["user_id"], template_id, cached_profiles)
        if not profile:
            continue

        exact_headers = profile.get("exact_headers", [])
        category = profile.get("category_name", "")
        src_row = prod.get("source_row", {})
        asin = prod.get("asin", "")

        cells = _map_row_to_template(
            src_row=src_row,
            exact_headers=exact_headers,
            mappings=mappings,
            fixed_values=_fixed_values_for_category(fv_entries, category),
            manual_edits=manual_edits,
            asin=asin,
            category=category,
        )

        rows.append({
            "asin": asin,
            "sku": prod.get("sku"),
            "template_id": template_id,
            "cells": {h: {"value": c.value, "status": c.status} for h, c in cells.items()},
        })

    return {"run_id": run_id, "rows": rows}


@router.patch("/{run_id}/values")
async def patch_values(run_id: str, body: ManualEditsRequest, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    existing = state.get("manual_edits", {})
    existing.update(body.edits)

    # If the caller supplied product hints (from Content page), inject any ASINs
    # that aren't already in resolved_products so the Preview can render them.
    if body.products:
        resolved = state.get("resolved_products", [])
        existing_asins = {p["asin"] for p in resolved}
        # Find the active template to attach to injected rows
        uploaded = state.get("uploaded_templates", [])
        tpl_id = uploaded[0] if uploaded else None
        tpl_filename = None
        if tpl_id:
            try:
                reg = TemplateRegistry()
                tpl = reg.get(tpl_id)
                tpl_filename = tpl.filename if tpl else None
            except Exception:
                pass

        for hint in body.products:
            if hint.asin not in existing_asins:
                resolved.append({
                    "asin": hint.asin,
                    "sku": hint.sku,
                    "product_name": hint.product_name,
                    "template_id": tpl_id,
                    "template_filename": tpl_filename,
                    "source_id": "content_converter",
                    "source_row": {},
                    "status": "matched",
                })
                existing_asins.add(hint.asin)

        run_store.update_run(current_user["user_id"], run_id, manual_edits=existing, resolved_products=resolved)
    else:
        run_store.update_run(current_user["user_id"], run_id, manual_edits=existing)

    return {"run_id": run_id, "edit_count": len(existing)}


@router.post("/{run_id}/sot-fill")
async def sot_fill_run(run_id: str, current_user: dict = Depends(get_current_user)):
    """Pull all SOT fields (EAN, weight, images, etc.) for every product in the run."""
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    products = state.get("resolved_products", [])
    if not products:
        return {"filled": 0, "total": 0}

    asins = [p["asin"] for p in products if p.get("asin")]

    try:
        import yaml
        from app.sources.configurator_sot_source import batch_get_sot_rows
        sot_mappings_path = _CONFIG_DIR / "configurator_sot_mappings.yaml"
        sot_mappings: list[dict] = []
        if sot_mappings_path.exists():
            with open(sot_mappings_path, encoding="utf-8") as _f:
                sot_mappings = yaml.safe_load(_f).get("mappings", [])
        sot_rows_by_asin, _ = batch_get_sot_rows(asins, sot_mappings)
    except Exception as exc:
        logger.error("sot_fill_run: batch_get_sot_rows failed: %s", exc)
        raise HTTPException(status_code=503, detail=f"SOT fill failed: {exc}")

    uploaded_templates = state.get("uploaded_templates", [])
    default_tpl_id = uploaded_templates[0] if uploaded_templates else None
    default_tpl_filename = None
    if default_tpl_id:
        try:
            reg = _load_registry(current_user["user_id"])
            tpl = reg.find_by_id(default_tpl_id)
            default_tpl_filename = tpl.filename if tpl else None
        except Exception:
            pass

    filled = 0
    manual_edits = state.get("manual_edits", {})
    edits_changed = False

    for prod in products:
        asin = prod.get("asin", "")
        sot_row = sot_rows_by_asin.get(asin, {})
        if not sot_row:
            continue
        # SOT fills gaps — existing source_row values take priority
        existing_row = prod.get("source_row") or {}
        merged = {**sot_row, **{k: v for k, v in existing_row.items() if v}}

        # Brand/MPN are fixed by marketplace — always overwrite, never from source
        sku = sot_row.get("Supplier Part Number") or prod.get("sku", "")
        if sot_row.get("__is_us__") == "true":
            brand, mpn = "Neighbour Market", f"NM-{sku}"
            # Blank out UK document fields for US ASINs
            for _i in range(1, 4):
                for _f in (
                    f"Document File Name or URL {_i}",
                    f"Document Type {_i}",
                    f"Legal Document Type {_i}",
                    f"Document Region {_i}",
                ):
                    manual_edits[f"{asin}::{_f}"] = ""
        else:
            brand, mpn = "LEDSone", f"WF-{sku}"
        merged["Brand"] = brand
        merged["Manufacturer Part Number"] = mpn
        manual_edits[f"{asin}::Brand"] = brand
        manual_edits[f"{asin}::Manufacturer Part Number"] = mpn
        edits_changed = True

        prod["source_row"] = merged
        # Assign template if not already set
        if not prod.get("template_id") and default_tpl_id:
            prod["template_id"] = default_tpl_id
            prod["template_filename"] = default_tpl_filename
        # Update SKU from SOT
        spn = sot_row.get("Supplier Part Number")
        if spn:
            prod["sku"] = spn
            # Also write directly into manual_edits so preview always shows it
            edit_key = f"{asin}::Supplier Part Number"
            if not manual_edits.get(edit_key):
                manual_edits[edit_key] = spn
                edits_changed = True
        prod["source_id"] = prod.get("source_id") or "configurator_sot"
        prod["status"] = "matched"
        filled += 1

    run_store.update_run(current_user["user_id"], run_id, resolved_products=products,
                         **({"manual_edits": manual_edits} if edits_changed else {}))
    return {"filled": filled, "total": len(products)}


@router.post("/{run_id}/validate")
async def validate_run(run_id: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    products = state.get("resolved_products", [])
    manual_edits = state.get("manual_edits", {})
    cached_profiles = state.get("cached_profiles", {})
    mappings = _load_field_mappings()
    fv_entries = _load_fixed_values()

    required_fields: set[str] = set()
    for m in mappings:
        if m.get("required"):
            required_fields.add(m.get("wayfair_field", ""))

    findings: list[dict] = []

    for prod in products:
        template_id = prod.get("template_id")
        asin = prod.get("asin", "")
        sku = prod.get("sku")

        if not template_id:
            findings.append({
                "asin": asin,
                "sku": sku,
                "column": "template",
                "value": "",
                "severity": "blocking",
                "error_type": "no_template",
                "recommended_action": "Assign a template to this product",
            })
            continue

        if prod.get("status") == "unmatched":
            findings.append({
                "asin": asin,
                "sku": sku,
                "column": "source_data",
                "value": "",
                "severity": "error",
                "error_type": "unmatched_asin",
                "recommended_action": "Provide source data for this ASIN",
            })
            continue

        profile = _load_profile(current_user["user_id"], template_id, cached_profiles)
        if not profile:
            continue

        exact_headers = profile.get("exact_headers", [])
        category = profile.get("category_name", "")
        src_row = prod.get("source_row", {})

        # Warn if Base Cost was derived from an unverified (wrong_sku=1) listing
        if src_row.get("__base_cost_unverified__") == "true":
            findings.append({
                "asin": asin,
                "sku": sku,
                "column": "Base Cost",
                "value": src_row.get("Base Cost", ""),
                "severity": "warning",
                "error_type": "unverified_base_cost",
                "recommended_action": "Base Cost calculated from a listing flagged as wrong_sku. Please verify the SKU in listing_data.",
            })

        cells = _map_row_to_template(
            src_row=src_row,
            exact_headers=exact_headers,
            mappings=mappings,
            fixed_values=_fixed_values_for_category(fv_entries, category),
            manual_edits=manual_edits,
            asin=asin,
            category=category,
        )

        for header, cell in cells.items():
            if cell.status == "missing" and header in required_fields:
                findings.append({
                    "asin": asin,
                    "sku": sku,
                    "column": header,
                    "value": "",
                    "severity": "blocking",
                    "error_type": "missing_required",
                    "recommended_action": f"Provide a value for required field: {header}",
                })

    run_store.update_run(current_user["user_id"], run_id, validation_findings=findings, status="validated")
    return {"run_id": run_id, "finding_count": len(findings), "findings": findings}


@router.get("/{run_id}/findings")
async def get_findings(run_id: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    raw = state.get("validation_findings", [])
    grouped: dict[str, list] = {"blocking": [], "error": [], "warning": [], "review": [], "info": []}
    for f in raw:
        sev = f.get("severity", "info")
        grouped.setdefault(sev, []).append(f)

    return grouped


@router.post("/{run_id}/export")
async def export_run(run_id: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    products = state.get("resolved_products", [])
    manual_edits = state.get("manual_edits", {})
    cached_profiles = state.get("cached_profiles", {})
    mappings = _load_field_mappings()
    fv_entries = _load_fixed_values()
    registry = _load_registry(current_user["user_id"])

    export_dir = run_store.export_dir(current_user["user_id"], run_id)
    export_dir.mkdir(parents=True, exist_ok=True)

    # Group products by template_id
    by_template: dict[str, list[dict]] = {}
    for prod in products:
        tid = prod.get("template_id")
        if tid:
            by_template.setdefault(tid, []).append(prod)

    export_files = []

    for template_id, prods in by_template.items():
        entry = registry.find_by_id(template_id)

        profile = _load_profile(current_user["user_id"], template_id, cached_profiles)
        if not profile:
            continue

        exact_headers = profile.get("exact_headers", [])
        category = profile.get("category_name", "unknown")
        listing_sheet = profile.get("listing_sheet")
        header_row = profile.get("listing_header_row") or 4
        raw_write_start = (entry.safe_write_start_row if entry else None) or profile.get("safe_write_start_row") or (header_row + 1)
        # Clamp: if the registry stored a sample-data row (e.g. 508) use a
        # sensible fallback. Wayfair templates have ≤ 6 instruction rows below
        # the header, so anything beyond header_row + 10 is stale sample data.
        safe_write_start = raw_write_start if raw_write_start <= header_row + 10 else header_row + 4

        # Find the original template file
        entry_filename = entry.filename if entry else Path(profile.get("source_file", "")).name
        template_path = Path("outputs/uploads") / current_user["user_id"] / entry_filename if entry_filename else None
        if template_path and not template_path.exists():
            alt = Path(profile.get("source_file", "")) if profile.get("source_file") else None
            template_path = alt if alt and alt.exists() else template_path

        # Build product data rows
        product_rows: list[dict[str, str]] = []
        # extra_image_rows: [{sku, image_url}] for additional images sheet
        extra_image_rows: list[dict[str, str]] = []
        sku_col_header = "Supplier Part Number"
        img_col_header = "Image File Name or URL"
        for prod in prods:
            asin = prod.get("asin", "")
            src_row = prod.get("source_row", {})
            cells = _map_row_to_template(
                src_row=src_row,
                exact_headers=exact_headers,
                mappings=mappings,
                fixed_values=_fixed_values_for_category(fv_entries, category),
                manual_edits=manual_edits,
                asin=asin,
                category=category,
            )
            row_dict = {h: c.value for h, c in cells.items()}
            # Collect extra images (stored as newline-separated in __extra_images__)
            extra_raw = src_row.get("__extra_images__", "")
            sku_val = src_row.get("Supplier Part Number") or src_row.get("__sku__") or prod.get("sku") or ""
            for img_url in (extra_raw.split("\n") if extra_raw else []):
                if img_url.strip():
                    extra_image_rows.append({sku_col_header: sku_val, img_col_header: img_url.strip()})
            product_rows.append(row_dict)

        safe_cat = re.sub(r"[^\w\-]", "_", category)

        # Export as CSV
        import csv as csv_module
        csv_filename = f"{safe_cat}_export.csv"
        csv_path = export_dir / csv_filename
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv_module.DictWriter(f, fieldnames=exact_headers)
            writer.writeheader()
            writer.writerows(product_rows)

        export_files.append({
            "filename": csv_filename,
            "template_id": template_id,
            "file_type": "csv",
            "product_count": len(prods),
        })

        # Export as XLSX (copy template if available, else create fresh)
        xlsx_filename = f"{safe_cat}_export.xlsx"
        xlsx_path = export_dir / xlsx_filename

        if template_path and template_path.exists() and listing_sheet:
            try:
                wb = openpyxl.load_workbook(str(template_path))
                if listing_sheet in wb.sheetnames:
                    ws = wb[listing_sheet]
                else:
                    ws = wb.active

                # Build header->col map
                header_col_map: dict[str, int] = {}
                for c in range(1, (ws.max_column or 1) + 1):
                    v = ws.cell(header_row, c).value
                    if v is not None:
                        header_col_map[str(v).strip()] = c

                # Clear only from safe_write_start_row onwards — preserves
                # instruction/definition rows between header and data area.
                max_col = ws.max_column or 1
                for r in range(safe_write_start, (ws.max_row or safe_write_start) + 1):
                    for c in range(1, max_col + 1):
                        ws.cell(r, c).value = None

                for row_offset, row_data in enumerate(product_rows):
                    write_row = safe_write_start + row_offset
                    for header, value in row_data.items():
                        col = header_col_map.get(header)
                        if col and value:
                            ws.cell(write_row, col).value = value

                # Write additional images sheet if present
                if extra_image_rows:
                    img_sheet_names = profile.get("image_sheets", [])
                    img_sheet = next(
                        (n for n in img_sheet_names if n in wb.sheetnames), None
                    )
                    if img_sheet:
                        ws_img = wb[img_sheet]
                        # Find the header row for Supplier Part Number and image URL
                        img_sku_col, img_url_col = None, None
                        for r in range(1, min(10, ws_img.max_row or 1) + 1):
                            for c in range(1, (ws_img.max_column or 1) + 1):
                                v = str(ws_img.cell(r, c).value or "").strip()
                                if v == sku_col_header:
                                    img_sku_col = c
                                elif v == img_col_header:
                                    img_url_col = c
                            if img_sku_col and img_url_col:
                                img_write_start = r + 1
                                break
                        if img_sku_col and img_url_col:
                            for ri, erow in enumerate(extra_image_rows):
                                ws_img.cell(img_write_start + ri, img_sku_col).value = erow[sku_col_header]
                                ws_img.cell(img_write_start + ri, img_url_col).value = erow[img_col_header]

                wb.save(str(xlsx_path))
                wb.close()
            except Exception as exc:
                # Fall back: create fresh
                _write_fresh_xlsx(xlsx_path, exact_headers, product_rows)
        else:
            _write_fresh_xlsx(xlsx_path, exact_headers, product_rows)

        export_files.append({
            "filename": xlsx_filename,
            "template_id": template_id,
            "file_type": "xlsx",
            "product_count": len(prods),
        })

    run_store.update_run(current_user["user_id"], run_id, export_files=export_files, status="exported")
    return ExportResponse(run_id=run_id, export_files=export_files)


def _write_fresh_xlsx(path: Path, headers: list[str], rows: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(headers, 1):
        ws.cell(1, c).value = h
    for r, row in enumerate(rows, 2):
        for c, h in enumerate(headers, 1):
            v = row.get(h, "")
            if v:
                ws.cell(r, c).value = v
    wb.save(str(path))
    wb.close()


@router.get("/{run_id}/downloads")
async def list_downloads(run_id: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    export_dir = run_store.export_dir(current_user["user_id"], run_id)
    files = []
    export_meta = {f["filename"]: f for f in state.get("export_files", [])}

    if export_dir.exists():
        for fp in sorted(export_dir.iterdir()):
            meta = export_meta.get(fp.name, {})
            files.append({
                "filename": fp.name,
                "size_bytes": fp.stat().st_size,
                "template_id": meta.get("template_id"),
                "file_type": meta.get("file_type", fp.suffix.lstrip(".")),
            })

    return files


@router.get("/{run_id}/downloads/{filename}")
async def download_file(run_id: str, filename: str, current_user: dict = Depends(get_current_user)):
    state = run_store.get_run(current_user["user_id"], run_id)
    if not state:
        raise HTTPException(status_code=404, detail=f"Run {run_id!r} not found")

    # Sanitize filename
    safe_filename = Path(filename).name
    file_path = run_store.export_dir(current_user["user_id"], run_id) / safe_filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail=f"File {filename!r} not found")

    media_type = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if safe_filename.endswith(".xlsx")
        else "text/csv; charset=utf-8"
    )
    return FileResponse(path=str(file_path), filename=safe_filename, media_type=media_type)
