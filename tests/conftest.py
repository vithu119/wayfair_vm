"""Shared fixtures for unit tests using synthetic openpyxl workbooks."""
from __future__ import annotations
import pytest
import openpyxl
import io
import tempfile
from pathlib import Path


def make_standard_wb(headers: list[str], required_row: list[str] | None = None) -> openpyxl.Workbook:
    """Create an in-memory standard-format Wayfair listing workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "9999 - Test Category"

    # Row 1: internal keys
    for c, h in enumerate(headers, 1):
        ws.cell(1, c).value = f"core::{h.replace(' ', '')}"
    # Row 2: group
    for c in range(1, len(headers) + 1):
        ws.cell(2, c).value = "Basic Information"
    # Row 3: required status
    req = required_row or ["Required"] * len(headers)
    for c, r in enumerate(req, 1):
        ws.cell(3, c).value = r
    # Row 4: human-readable headers  ← THE display headers
    for c, h in enumerate(headers, 1):
        ws.cell(4, c).value = h
    # Row 5: instructions
    for c, h in enumerate(headers, 1):
        ws.cell(5, c).value = f"Enter {h}"
    # Row 6: data type
    for c in range(1, len(headers) + 1):
        ws.cell(6, c).value = "Text"
    # Row 7: default
    for c in range(1, len(headers) + 1):
        ws.cell(7, c).value = "N/A"

    return wb


def save_wb_to_temp(wb: openpyxl.Workbook, suffix: str = ".xlsx") -> Path:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    wb.save(tmp.name)
    tmp.close()
    return Path(tmp.name)
