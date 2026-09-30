"""Unit tests for valid-value extraction."""
import openpyxl
import pytest
from app.analyser.valid_values import extract_valid_values


def _make_vv_sheet(data: dict[str, list[str]]) -> openpyxl.worksheet.worksheet.Worksheet:
    wb = openpyxl.Workbook()
    ws = wb.active
    cols = list(data.keys())
    for c, col_name in enumerate(cols, 1):
        ws.cell(1, c).value = col_name
        for r, val in enumerate(data[col_name], 2):
            ws.cell(r, c).value = val
    return ws


class TestValidValueExtraction:
    def test_extracts_correct_field_names(self):
        ws = _make_vv_sheet({"Color": ["Red", "Blue"], "Size": ["Small", "Large"]})
        tables = extract_valid_values(ws, "Valid Values", max_col=2, max_row=3)
        names = {t.field_name for t in tables}
        assert "Color" in names
        assert "Size" in names

    def test_extracts_exact_values(self):
        ws = _make_vv_sheet({"Finish": ["Brushed Nickel", "Oil Rubbed Bronze", "Chrome"]})
        tables = extract_valid_values(ws, "Valid Values", max_col=1, max_row=4)
        assert len(tables) == 1
        assert tables[0].values == ["Brushed Nickel", "Oil Rubbed Bronze", "Chrome"]

    def test_preserves_exact_capitalization(self):
        ws = _make_vv_sheet({"Type": ["LED", "CFL", "Incandescent", "halogen"]})
        tables = extract_valid_values(ws, "Valid Values", max_col=1, max_row=5)
        assert "LED" in tables[0].values
        assert "halogen" in tables[0].values

    def test_cell_range_recorded(self):
        ws = _make_vv_sheet({"Color": ["Red", "Blue", "Green"]})
        tables = extract_valid_values(ws, "Valid Values", max_col=1, max_row=4)
        assert tables[0].cell_range  # non-empty
        assert "A" in tables[0].cell_range

    def test_empty_sheet_returns_empty(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        tables = extract_valid_values(ws, "Valid Values", max_col=0, max_row=0)
        assert tables == []
